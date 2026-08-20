#!/usr/bin/env python3
"""Test end-to-end ilościówki: surowy plik → nasz Ip → % czystości vs Excel lab.

Dopasowuje pliki z „TPrA - pozytywy" do wierszy „Excel do komercji.xlsx"
po nazwie próbki i dodanej objętości zapisanej w nazwie pliku
(np. 5.2_1200ul_TPrA.txt ↔ próbka 5.2, V = 1200 µL).

To trzecia, niezależna ścieżka weryfikacji: poprzednie porównywały nasze Ip
z Ip lab (CC) oraz nasze % z % lab przy TYM SAMYM prądzie. Tutaj cały łańcuch
liczymy od surowego woltamperogramu.
"""
import glob
import os
import re

import numpy as np

from test_ities_local import load_algorithm

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
POZ_DIR = os.path.join(ROOT, "06_ities_update_20260727", "TPrA - pozytywy")
XLSX = os.path.join(ROOT, "06_ities_update_20260727", "Excel do komercji.xlsx")


def wiersze_excela():
    import openpyxl
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Sheet1"]
    out = {}
    for r in range(8, 50):
        nazwa, masa, vmax, prad, proc = (ws[f"B{r}"].value, ws[f"C{r}"].value,
                                         ws[f"H{r}"].value, ws[f"S{r}"].value,
                                         ws[f"V{r}"].value)
        if nazwa is None or masa is None or vmax is None:
            continue
        if not isinstance(prad, (int, float)) or not isinstance(proc, (int, float)):
            continue
        out[str(nazwa).strip()] = {"masa_mg": masa, "v_dodane_uL": round(vmax * 1e6),
                                   "ip_lab_uA": prad * 1e6, "pct_lab": proc}
    return out


def dopasuj_pliki(probki):
    """Pliki, których nazwa zaczyna się od nazwy próbki i zawiera właściwą objętość."""
    pary = []
    for path in sorted(glob.glob(os.path.join(POZ_DIR, "*"))):
        base = os.path.basename(path)
        if base.startswith(".") or not os.path.isfile(path):
            continue
        m = re.match(r"([0-9]+\.[0-9]+|[0-9]+-[0-9]+[AB]?)[_ ]", base)
        if not m:
            continue
        nazwa = m.group(1)
        if nazwa not in probki:
            continue
        mv = re.search(r"_(\d+)\s*ul", base, re.I)
        if not mv or int(mv.group(1)) != probki[nazwa]["v_dodane_uL"]:
            continue
        pary.append((nazwa, path))
    return pary


def main():
    algo = load_algorithm()
    probki = wiersze_excela()
    pary = dopasuj_pliki(probki)
    print(f"Dopasowanych par plik↔próbka: {len(pary)}\n")
    print(f"{'próbka':>8} {'plik':<34} {'Ip nasz':>8} {'Ip lab':>8} {'Δ Ip %':>7} "
          f"{'% nasz':>8} {'% lab':>8} {'Δ p.p.':>8}")
    dip, dpct = [], []
    for nazwa, path in pary:
        p = probki[nazwa]
        algo.RESULTS.clear()
        algo.ustaw_probke(os.path.basename(path), p["masa_mg"], p["v_dodane_uL"])
        try:
            r = algo.analyze(os.path.basename(path), open(path, "rb").read())
        except Exception as e:                                    # noqa: BLE001
            print(f"{nazwa:>8} {os.path.basename(path)[:34]:<34}  CRASH: {e}")
            continue
        ip = r.get("Ip_analyte_fwd_uA")
        pct = r.get("purity_pct")
        if ip is None:
            print(f"{nazwa:>8} {os.path.basename(path)[:34]:<34}  brak Ip "
                  f"(status {r.get('status')})")
            continue
        d_ip = 100 * (ip - p["ip_lab_uA"]) / p["ip_lab_uA"]
        d_pct = (pct - p["pct_lab"]) if pct is not None else float("nan")
        dip.append(d_ip)
        dpct.append(d_pct)
        print(f"{nazwa:>8} {os.path.basename(path)[:34]:<34} {ip:8.3f} "
              f"{p['ip_lab_uA']:8.3f} {d_ip:7.1f} {pct:8.4f} {p['pct_lab']:8.4f} {d_pct:8.4f}")

    if dip:
        print(f"\nΔ Ip: mediana {np.median(dip):+.1f} %, |Δ| ≤ 20 % dla "
              f"{sum(abs(d) <= 20 for d in dip)}/{len(dip)} par")
        print(f"Δ czystości: mediana {np.median(dpct):+.4f} p.p., "
              f"maks. |Δ| {max(abs(d) for d in dpct):.4f} p.p.")


if __name__ == "__main__":
    main()
